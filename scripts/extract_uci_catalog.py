"""Build a verified UCI StockCode->Description mapping, without re-training.

Reads all worksheets of Online Retail II in streaming mode, supports CSV and
XLSX. Writes only StockCodes used by models/uci_rules.csv.

Windows (PowerShell, single line):
  python scripts/extract_uci_catalog.py --input "C:\\path\\online_retail_II.xlsx"

Kaggle (after the master notebook exported rules):
  python scripts/extract_uci_catalog.py --input /kaggle/input/datasets/wafabila/bundlerecommender-dataset/onlineretailii/online_retail_II.xlsx --rules /kaggle/working/bundlerecommender_master_v1/models/uci_rules.csv --output /kaggle/working/uci_product_catalog.csv
"""
from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterator


def normalize_header(value: object) -> str:
    return "".join(ch for ch in str(value or "").lower() if ch.isalnum())


def source_records(path: Path) -> Iterator[tuple[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Install Excel reader: python -m pip install openpyxl") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                iterator = ws.iter_rows(values_only=True)
                headers = next(iterator, None)
                if not headers:
                    continue
                normalized = [normalize_header(h) for h in headers]
                if "stockcode" not in normalized or "description" not in normalized:
                    print(f"SKIP worksheet '{ws.title}': no StockCode/Description columns")
                    continue
                code_idx = normalized.index("stockcode")
                desc_idx = normalized.index("description")
                print(f"Reading worksheet '{ws.title}'...")
                for row in iterator:
                    if len(row) > max(code_idx, desc_idx):
                        yield str(row[code_idx] or "").strip().upper(), str(row[desc_idx] or "").strip()
        finally:
            wb.close()
    elif suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="", errors="replace") as stream:
            reader = csv.DictReader(stream)
            headers = {normalize_header(h): h for h in (reader.fieldnames or [])}
            if not {"stockcode", "description"}.issubset(headers):
                raise ValueError(f"Source CSV requires StockCode and Description. Actual: {reader.fieldnames}")
            for row in reader:
                yield (row.get(headers["stockcode"]) or "").strip().upper(), (row.get(headers["description"]) or "").strip()
    else:
        raise ValueError("Use the original .xlsx or a transaction .csv containing StockCode and Description")


def required_skus(rule_file: Path) -> set[str]:
    with rule_file.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        fields = set(reader.fieldnames or [])
        if not {"antecedent_sku", "consequent_sku"}.issubset(fields):
            raise ValueError(f"Invalid UCI rule schema: {sorted(fields)}")
        return {str(row[column]).strip().upper() for row in reader
                for column in ("antecedent_sku", "consequent_sku") if row.get(column)}


def create_catalog(source: Path, rules: Path, output: Path) -> dict:
    if not source.is_file():
        raise FileNotFoundError(f"Original UCI source not found: {source}")
    if not rules.is_file():
        raise FileNotFoundError(f"Rule artifact not found: {rules}")
    wanted = required_skus(rules)
    counts: dict[str, Counter[str]] = defaultdict(Counter)
    seen = 0
    for sku, description in source_records(source):
        seen += 1
        clean = " ".join(description.split())
        if sku in wanted and clean and clean.casefold() not in {"nan", "none", "null"}:
            counts[sku][clean] += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["stock_code", "description"])
        writer.writeheader()
        for sku in sorted(wanted):
            if counts[sku]:
                chosen = min(counts[sku], key=lambda d: (-counts[sku][d], d))
                writer.writerow({"stock_code": sku, "description": chosen})
    return {"source_rows_scanned": seen, "rule_skus": len(wanted),
            "named_skus": sum(bool(counts[s]) for s in wanted),
            "unnamed_skus": sum(not bool(counts[s]) for s in wanted), "output": str(output)}


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract verified UCI descriptions (no model training)")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--rules", type=Path, default=Path("models/uci_rules.csv"))
    parser.add_argument("--output", type=Path, default=Path("models/uci_product_catalog.csv"))
    args = parser.parse_args()
    for key, value in create_catalog(args.input, args.rules, args.output).items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()
